from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from bale_inviter.domain.models import ParsedContactRow

NAME_HEADERS = {"name", "نام", "اسم", "fullname", "full_name", "full name", "نام و نام خانوادگی"}
PHONE_HEADERS = {
    "phone",
    "mobile",
    "tel",
    "شماره",
    "تلفن",
    "موبایل",
    "شماره موبایل",
    "phone_number",
    "phonenumber",
    "phone number",
}


class ImportParseError(ValueError):
    pass


def parse_contacts_file(path: str | Path) -> list[ParsedContactRow]:
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"Import file not found: {file_path}")
    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(file_path)
    if suffix == ".csv":
        return _parse_csv(file_path)
    raise ImportParseError(f"Unsupported file type: {suffix}. Use .csv or .xlsx")


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower()


def _map_headers(headers: list[str]) -> tuple[int, int]:
    name_idx = phone_idx = None
    for index, header in enumerate(headers):
        key = _normalize_header(header)
        if key in NAME_HEADERS and name_idx is None:
            name_idx = index
        if key in PHONE_HEADERS and phone_idx is None:
            phone_idx = index
    if name_idx is None and phone_idx is None and len(headers) >= 2:
        return 0, 1
    if name_idx is None or phone_idx is None:
        raise ImportParseError(
            "Could not detect name/phone columns. Expected headers like name, phone, نام, شماره."
        )
    return name_idx, phone_idx


def _row_to_parsed(
    row_number: int,
    values: list[object],
    headers: list[str],
    name_idx: int,
    phone_idx: int,
) -> ParsedContactRow | None:
    if not any(str(v).strip() for v in values if v is not None):
        return None
    name = str(values[name_idx] if name_idx < len(values) else "").strip()
    phone = str(values[phone_idx] if phone_idx < len(values) else "").strip()
    extra: dict[str, str] = {}
    for idx, header in enumerate(headers):
        if idx in {name_idx, phone_idx} or not header:
            continue
        if idx < len(values) and values[idx] not in (None, ""):
            extra[header] = str(values[idx])
    return ParsedContactRow(row_number=row_number, name=name, phone=phone, extra=extra)


def _parse_csv(path: Path) -> list[ParsedContactRow]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1256"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(4096)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(handle, dialect)
                rows = list(reader)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
            rows = []
    else:
        raise ImportParseError(f"Could not decode CSV: {last_error}")

    if not rows:
        return []
    headers = [_normalize_header(item) for item in rows[0]]
    name_idx, phone_idx = _map_headers(headers)
    parsed: list[ParsedContactRow] = []
    for offset, raw in enumerate(rows[1:], start=2):
        item = _row_to_parsed(offset, list(raw), headers, name_idx, phone_idx)
        if item is not None:
            parsed.append(item)
    return parsed


def _parse_xlsx(path: Path) -> list[ParsedContactRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [_normalize_header(item) for item in rows[0]]
    name_idx, phone_idx = _map_headers(headers)
    parsed: list[ParsedContactRow] = []
    for offset, raw in enumerate(rows[1:], start=2):
        item = _row_to_parsed(offset, list(raw or ()), headers, name_idx, phone_idx)
        if item is not None:
            parsed.append(item)
    return parsed
